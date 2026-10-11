#!/usr/bin/env python3
"""
build_tutorial_pptx.py  --  Self-demonstrating tutorial slide deck for
building scientific presentations with python-pptx.

Every technique documented in beamer/README.md is both explained and
*demonstrated live* on the slides themselves: tables, matplotlib figures,
LaTeX-rendered equations, mixed formatting, hyperlinks, and slide numbers.

Usage
-----
    python beamer/build_tutorial_pptx.py
    python beamer/build_tutorial_pptx.py -o my_tutorial.pptx

Requires: python-pptx, matplotlib, numpy
"""

from __future__ import annotations

import argparse
import io
import uuid
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_OUTPUT = REPO_ROOT / "beamer" / "python_pptx_tutorial.pptx"
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# Colour palette
CLR_DARK      = RGBColor(0x23, 0x37, 0x3B)
CLR_TEAL      = RGBColor(0x00, 0x77, 0x88)
CLR_ACCENT    = RGBColor(0xEB, 0x81, 0x1E)
CLR_TUMOR     = RGBColor(0xCC, 0x33, 0x33)
CLR_CAVITY    = RGBColor(0x33, 0x66, 0xCC)
CLR_WM        = RGBColor(0x33, 0x99, 0x66)
CLR_WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
CLR_LIGHT_BG  = RGBColor(0xF5, 0xF5, 0xF0)
CLR_MED_GRAY  = RGBColor(0x88, 0x88, 0x88)
CLR_DARK_TEXT  = RGBColor(0x2D, 0x2D, 0x2D)
CLR_CODE_BG   = RGBColor(0x2B, 0x2B, 0x2B)
CLR_CODE_FG   = RGBColor(0xD4, 0xD4, 0xD4)
CLR_KW_BLUE   = RGBColor(0x56, 0x9C, 0xD6)
CLR_STR_BROWN = RGBColor(0xCE, 0x91, 0x78)
CLR_COMMENT   = RGBColor(0x6A, 0x99, 0x55)


# ---------------------------------------------------------------------------
# Core helpers (same API as build_pptx.py)
# ---------------------------------------------------------------------------

def _bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _tb(slide, l, t, w, h, text, *, sz=18, bold=False, color=CLR_DARK_TEXT,
        align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(sz)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font
    p.alignment = align
    return box


def _rich(slide, l, t, w, h, paras, *, sz=18, color=CLR_DARK_TEXT,
          font="Calibri", align=PP_ALIGN.LEFT):
    """paras: list of (text, bold, size_override|None, color_override|None)."""
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(paras):
        txt, bld = item[0], item[1]
        sz_o = item[2] if len(item) > 2 else None
        clr_o = item[3] if len(item) > 3 else None
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(4)
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(sz_o or sz)
        r.font.bold = bld
        r.font.color.rgb = clr_o or color
        r.font.name = font
    return box


def _notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def _title_bar(slide, text):
    box = _tb(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.9),
              f"  {text}", sz=26, bold=True, color=CLR_WHITE,
              font="Calibri Light")
    box.fill.solid()
    box.fill.fore_color.rgb = CLR_TEAL
    return box


def _slide_num(slide):
    box = slide.shapes.add_textbox(
        Inches(12.4), Inches(7.0), Inches(0.8), Inches(0.4))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    fld = etree.SubElement(
        p._p, qn("a:fld"),
        attrib={"type": "slidenum",
                "id": "{" + str(uuid.uuid4()).upper() + "}"})
    rPr = etree.SubElement(fld, qn("a:rPr"), lang="en-US", dirty="0")
    rPr.set("sz", "1200")
    srgb = etree.SubElement(
        etree.SubElement(rPr, qn("a:solidFill")), qn("a:srgbClr"),
        val="888888")
    t = etree.SubElement(fld, qn("a:t"))
    t.text = "\u2039#\u203a"
    etree.SubElement(p._p, qn("a:endParaRPr"), lang="en-US", dirty="0")


# ---------------------------------------------------------------------------
# Tables helper
# ---------------------------------------------------------------------------

def _add_table(slide, l, t, w, h, data, *,
               col_widths=None, header_color=None, font_size=14,
               alt_row_color=None):
    rows, cols = len(data), len(data[0])
    shape = slide.shapes.add_table(rows, cols, l, t, w, h)
    tbl = shape.table
    if col_widths:
        for j, cw in enumerate(col_widths):
            tbl.columns[j].width = cw
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = str(val)
            for para in cell.text_frame.paragraphs:
                para.font.size = Pt(font_size)
                para.font.name = "Calibri"
                para.alignment = PP_ALIGN.CENTER
            if i == 0:
                for para in cell.text_frame.paragraphs:
                    para.font.bold = True
                    para.font.color.rgb = CLR_WHITE
                if header_color:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = header_color
            elif alt_row_color and i % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = alt_row_color
    return shape


def _set_cell_border(cell, color="007788", width="12700"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        ln = etree.SubElement(tcPr, ns + edge, w=width)
        sf = etree.SubElement(ln, ns + "solidFill")
        etree.SubElement(sf, ns + "srgbClr", val=color)


# ---------------------------------------------------------------------------
# Matplotlib → slide
# ---------------------------------------------------------------------------

def _mpl_to_slide(slide, fig, l, t, w, h=None, *, dpi=200):
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight",
                facecolor="white", edgecolor="none")
    buf.seek(0)
    if h:
        slide.shapes.add_picture(buf, l, t, w, h)
    else:
        slide.shapes.add_picture(buf, l, t, width=w)
    buf.close()


# ---------------------------------------------------------------------------
# LaTeX → PNG → slide
# ---------------------------------------------------------------------------

def _latex_to_buf(latex, fontsize=28, dpi=300, color="black"):
    fig, ax = plt.subplots(figsize=(0.01, 0.01))
    ax.axis("off")
    fig.patch.set_alpha(0.0)
    txt = ax.text(0, 0, f"${latex}$", fontsize=fontsize, color=color,
                  ha="left", va="baseline", transform=ax.transAxes)
    fig.canvas.draw()
    bb = txt.get_window_extent(fig.canvas.get_renderer())
    bb = bb.expanded(1.15, 1.3)
    bbi = bb.transformed(fig.dpi_scale_trans.inverted())
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, bbox_inches=bbi,
                transparent=True)
    plt.close(fig)
    buf.seek(0)
    return buf


def _add_latex(slide, latex, l, t, w, **kw):
    buf = _latex_to_buf(latex, **kw)
    slide.shapes.add_picture(buf, l, t, width=w)
    buf.close()


def _add_latex_dark(slide, latex, l, t, w, **kw):
    kw.setdefault("color", "white")
    _add_latex(slide, latex, l, t, w, **kw)


# ---------------------------------------------------------------------------
# Code block helper
# ---------------------------------------------------------------------------

def _code_box(slide, l, t, w, h, code_text, *, sz=11):
    box = slide.shapes.add_textbox(l, t, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = CLR_CODE_BG
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(code_text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(1)
        p.space_before = Pt(0)
        r = p.add_run()
        r.text = line
        r.font.size = Pt(sz)
        r.font.name = "Menlo"
        r.font.color.rgb = CLR_CODE_FG
    return box


# ---------------------------------------------------------------------------
# Mixed-formatting demo helper
# ---------------------------------------------------------------------------

def _mixed_run(p, text, *, sz=20, bold=False, italic=False,
               color=CLR_DARK_TEXT, font="Calibri", underline=False,
               hyperlink=None):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(sz)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    r.font.underline = underline
    if hyperlink:
        r.hyperlink.address = hyperlink
    return r


# ═══════════════════════════════════════════════════════════════════════════
# SLIDES
# ═══════════════════════════════════════════════════════════════════════════

def slide_title(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, CLR_DARK)
    _tb(s, Inches(0.9), Inches(1.2), Inches(11), Inches(1.8),
        "Building Scientific Slide Decks\nwith python-pptx",
        sz=46, bold=True, color=CLR_WHITE, font="Calibri Light")
    _tb(s, Inches(0.9), Inches(3.5), Inches(11), Inches(1),
        "A self-demonstrating tutorial",
        sz=26, color=CLR_ACCENT, font="Calibri Light")
    _tb(s, Inches(0.9), Inches(5.0), Inches(11), Inches(1.5),
        "Glioma Brain Digital Twin project\n"
        "CAS \u2014 Mathematical Challenges in Brain Mechanics\n"
        "Finse, March 2026",
        sz=18, color=CLR_MED_GRAY, font="Calibri Light")
    _notes(s, "This deck is entirely generated by build_tutorial_pptx.py. "
              "Every technique it teaches is demonstrated live on the slides.")


def slide_overview(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, CLR_LIGHT_BG)
    _title_bar(s, "Tutorial Overview")
    _slide_num(s)

    _rich(s, Inches(0.8), Inches(1.3), Inches(5.5), Inches(5.5), [
        ("This deck covers:", True, 22),
        ("", False, 8),
        ("  1.  Quick start & architecture", False, 20),
        ("  2.  Tables  \u2014  _add_table(), borders, merging", False, 20),
        ("  3.  Figures  \u2014  PNGs, matplotlib, DPI tips", False, 20),
        ("  4.  Math: Approach A  \u2014  Unicode symbols", False, 20),
        ("  5.  Math: Approach B  \u2014  LaTeX \u2192 PNG", False, 20),
        ("  6.  Math: Approach C  \u2014  OOXML equations", False, 20),
        ("  7.  Mixed formatting & hyperlinks", False, 20),
        ("  8.  Slide numbering", False, 20),
        ("  9.  Workflow summary", False, 20),
        (" 10.  References", False, 20),
    ])

    _code_box(s, Inches(7.0), Inches(1.3), Inches(5.8), Inches(5.5),
              "# Generate this tutorial deck\n"
              "python beamer/build_tutorial_pptx.py\n"
              "\n"
              "# Generate the main talk deck\n"
              "python beamer/build_pptx.py\n"
              "\n"
              "# Custom output\n"
              "python beamer/build_pptx.py -o talk.pptx\n"
              "\n"
              "# Alternate figure directory\n"
              "python beamer/build_pptx.py \\\n"
              "    --fig-dir /path/to/figures", sz=14)

    _notes(s, "Both scripts live in beamer/ and share the same colour "
              "palette and helper functions.")


def slide_architecture(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, CLR_LIGHT_BG)
    _title_bar(s, "Architecture of build_pptx.py")
    _slide_num(s)

    _code_box(s, Inches(0.8), Inches(1.3), Inches(6.0), Inches(4.5),
              "build_pptx.py\n"
              "\u251c\u2500\u2500 Configuration\n"
              "\u2502   colours, fonts, slide 16:9\n"
              "\u251c\u2500\u2500 Helpers\n"
              "\u2502   \u251c\u2500 _set_slide_bg()\n"
              "\u2502   \u251c\u2500 _add_textbox()\n"
              "\u2502   \u251c\u2500 _add_rich_textbox()\n"
              "\u2502   \u251c\u2500 _add_image()\n"
              "\u2502   \u251c\u2500 _set_notes()\n"
              "\u2502   \u251c\u2500 _section_divider()\n"
              "\u2502   \u251c\u2500 _content_slide()\n"
              "\u2502   \u2514\u2500 _full_figure_slide()\n"
              "\u251c\u2500\u2500 Section builders\n"
              "\u2502   build_c01() \u2026 build_c09()\n"
              "\u2514\u2500\u2500 main() \u2192 CLI + build_deck()",
              sz=14)

    # Live colour palette table
    palette = [
        ["Variable", "Hex", "Use"],
        ["CLR_DARK",   "#23373B", "Dark backgrounds"],
        ["CLR_TEAL",   "#007788", "Title bars"],
        ["CLR_ACCENT", "#EB811E", "Orange accents"],
        ["CLR_TUMOR",  "#CC3333", "Tumor overlays"],
        ["CLR_CAVITY", "#3366CC", "Cavity overlays"],
        ["CLR_WM",     "#339966", "White-matter"],
    ]
    tbl_shape = _add_table(
        s, Inches(7.2), Inches(1.3), Inches(5.6), Inches(3.5),
        palette, header_color=CLR_TEAL, font_size=15,
        alt_row_color=RGBColor(0xE8, 0xF0, 0xF0))

    # Colour swatches
    swatch_colors = [CLR_DARK, CLR_TEAL, CLR_ACCENT,
                     CLR_TUMOR, CLR_CAVITY, CLR_WM]
    for i, c in enumerate(swatch_colors):
        sw = s.shapes.add_shape(
            1, Inches(7.25), Inches(1.85 + i * 0.50), Inches(0.35), Inches(0.35))
        sw.fill.solid()
        sw.fill.fore_color.rgb = c
        sw.line.fill.background()

    _tb(s, Inches(0.8), Inches(6.0), Inches(11), Inches(0.8),
        "Tip:  Keep colours in constants so every slide stays consistent.",
        sz=16, color=CLR_MED_GRAY, font="Calibri Light")

    _notes(s, "The colour palette is defined once as module-level constants "
              "and referenced everywhere. Matches preamble.tex for visual "
              "consistency between Beamer and PPTX.")


# ── TABLES ────────────────────────────────────────────────────────────────

def slide_tables_intro(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, CLR_DARK)
    _tb(s, Inches(0.9), Inches(1.5), Inches(11), Inches(1.2),
        "Section 1", sz=28, color=CLR_ACCENT, font="Calibri Light")
    _tb(s, Inches(0.9), Inches(2.5), Inches(11), Inches(2),
        "Tables", sz=48, bold=True, color=CLR_WHITE, font="Calibri Light")
    _tb(s, Inches(0.9), Inches(4.5), Inches(11), Inches(1.2),
        "_add_table()  \u00b7  borders via lxml  \u00b7  cell merging  "
        "\u00b7  alternating rows",
        sz=20, color=CLR_MED_GRAY, font="Calibri Light")
    _notes(s, "python-pptx has a full table API. The main gap is cell "
              "borders, which require lxml manipulation.")


def slide_tables_demo(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, CLR_LIGHT_BG)
    _title_bar(s, "Tables: MRI Acquisition Parameters")
    _slide_num(s)

    data = [
        ["Modality", "Resolution", "TR / TE", "Other"],
        ["T1w MPRAGE", "1.0 mm iso", "2300 / 2.98 ms", "176 slices"],
        ["DWI",        "2.5 mm iso", "7200 / 90 ms",   "b = 700 / 1000 / 2000"],
        ["BOLD",       "3.0 mm iso", "1700 / 30 ms",   "200 volumes"],
    ]
    tbl_shape = _add_table(
        s, Inches(0.8), Inches(1.3), Inches(7.0), Inches(2.8),
        data,
        col_widths=[Inches(1.6), Inches(1.6), Inches(1.8), Inches(2.0)],
        header_color=CLR_TEAL, font_size=16,
        alt_row_color=RGBColor(0xE8, 0xF4, 0xF0))

    # Apply custom borders to demonstrate the lxml technique
    tbl = tbl_shape.table
    for i in range(len(data)):
        for j in range(len(data[0])):
            _set_cell_border(tbl.cell(i, j), color="007788", width="9525")

    _code_box(s, Inches(8.3), Inches(1.3), Inches(4.6), Inches(4.2),
              'data = [\n'
              '  ["Modality","Resolution","TR/TE","Other"],\n'
              '  ["T1w",     "1.0mm iso", ...          ],\n'
              '  ...\n'
              ']\n'
              '_add_table(slide,\n'
              '  left, top, width, height,\n'
              '  data,\n'
              '  col_widths=[Inches(1.6), ...],\n'
              '  header_color=CLR_TEAL,\n'
              '  alt_row_color=RGBColor(...))\n'
              '\n'
              '# Borders (lxml workaround):\n'
              'for cell in all_cells:\n'
              '  _set_cell_border(cell,\n'
              '    color="007788", width="9525")',
              sz=12)

    _tb(s, Inches(0.8), Inches(4.4), Inches(7), Inches(0.5),
        "\u2191  Live table with teal header, alternating rows, and custom "
        "borders", sz=14, color=CLR_MED_GRAY)

    _notes(s, "This table is generated live by the script. The borders "
              "use the _set_cell_border() lxml workaround since python-pptx "
              "lacks a high-level border API.")


def slide_tables_advanced(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, CLR_LIGHT_BG)
    _title_bar(s, "Tables: Merging and Per-Cell Colour")
    _slide_num(s)

    data = [
        ["Tissue", "Shear modulus \u03bc (kPa)", "Fibre stiffness \u03ba"],
        ["Grey matter", "1.0", "0"],
        ["White matter", "1.5", "2.0"],
        ["Deep nuclei", "1.8", "0"],
        ["Tumor", "0.5 \u2013 2.0", "?"],
        ["CSF", "\u2248 0", "0"],
    ]
    tbl_shape = _add_table(
        s, Inches(0.8), Inches(1.3), Inches(6.5), Inches(3.6),
        data, header_color=CLR_TEAL, font_size=16,
        col_widths=[Inches(2.2), Inches(2.2), Inches(2.1)])

    tbl = tbl_shape.table
    # Highlight tumor row
    for j in range(3):
        cell = tbl.cell(4, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0xFF, 0xE0, 0xE0)
        for para in cell.text_frame.paragraphs:
            para.font.color.rgb = CLR_TUMOR

    _rich(s, Inches(7.8), Inches(1.3), Inches(5.0), Inches(5.0), [
        ("Key table techniques:", True, 20),
        ("", False, 8),
        ("  \u2022  Merge cells:", True, 18),
        ("    tbl.cell(0,0).merge(tbl.cell(0,2))", False, 14, CLR_TEAL),
        ("", False, 8),
        ("  \u2022  Per-cell colour:", True, 18),
        ("    cell.fill.solid()", False, 14, CLR_TEAL),
        ("    cell.fill.fore_color.rgb = CLR_TUMOR", False, 14, CLR_TEAL),
        ("", False, 8),
        ("  \u2022  Borders via lxml:", True, 18),
        ("    _set_cell_border(cell,", False, 14, CLR_TEAL),
        ('      color="CC3333", width="12700")', False, 14, CLR_TEAL),
        ("", False, 8),
        ("  \u2022  Alternating row shading:", True, 18),
        ("    alt_row_color=RGBColor(...)", False, 14, CLR_TEAL),
    ])

    _tb(s, Inches(0.8), Inches(5.2), Inches(6.5), Inches(0.5),
        "\u2191  Tumor row highlighted with per-cell fill + font colour",
        sz=14, color=CLR_MED_GRAY)

    _notes(s, "The tumor row uses per-cell fill and font colour to draw "
              "attention. Merge, borders, and banding are the three most "
              "commonly needed table features beyond the basics.")


# ── FIGURES ───────────────────────────────────────────────────────────────

def slide_figures_intro(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, CLR_DARK)
    _tb(s, Inches(0.9), Inches(1.5), Inches(11), Inches(1.2),
        "Section 2", sz=28, color=CLR_ACCENT, font="Calibri Light")
    _tb(s, Inches(0.9), Inches(2.5), Inches(11), Inches(2),
        "Figures and Images", sz=48, bold=True, color=CLR_WHITE,
        font="Calibri Light")
    _tb(s, Inches(0.9), Inches(4.5), Inches(11), Inches(1.2),
        "PNGs  \u00b7  matplotlib \u2192 slide  \u00b7  Freeview screenshots"
        "  \u00b7  DPI guide",
        sz=20, color=CLR_MED_GRAY, font="Calibri Light")


def slide_figures_mpl(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, CLR_LIGHT_BG)
    _title_bar(s, "Figures: Matplotlib \u2192 Slide (No Disk I/O)")
    _slide_num(s)

    # Generate a live connectivity-matrix-style figure
    np.random.seed(42)
    mat = np.corrcoef(np.random.randn(83, 200))
    np.fill_diagonal(mat, 0)

    fig, axes = plt.subplots(1, 2, figsize=(10, 4.5))
    im0 = axes[0].imshow(mat, cmap="RdBu_r", vmin=-0.3, vmax=0.3)
    axes[0].set_title("Simulated FC matrix (83\u00d783)", fontsize=12)
    plt.colorbar(im0, ax=axes[0], shrink=0.8)

    x = np.linspace(0, 30, 300)
    from math import factorial
    hrf = (x ** 5) * np.exp(-x) / 120 - 0.1 * (x ** 15) * np.exp(-x) / (
        factorial(15))
    hrf = hrf / np.max(hrf)
    axes[1].plot(x, hrf, color="#007788", linewidth=2)
    axes[1].set_xlabel("Time (s)", fontsize=11)
    axes[1].set_ylabel("Amplitude", fontsize=11)
    axes[1].set_title("Canonical HRF", fontsize=12)
    axes[1].axhline(0, color="grey", linewidth=0.5)
    fig.tight_layout()

    _mpl_to_slide(s, fig, Inches(0.6), Inches(1.2), Inches(7.5),
                  Inches(4.0), dpi=200)
    plt.close(fig)

    _code_box(s, Inches(8.5), Inches(1.2), Inches(4.4), Inches(5.0),
              'import io, matplotlib.pyplot as plt\n'
              '\n'
              'def _mpl_to_slide(slide, fig,\n'
              '                  l, t, w, h=None,\n'
              '                  *, dpi=200):\n'
              '    buf = io.BytesIO()\n'
              '    fig.savefig(buf, format="png",\n'
              '      dpi=dpi, bbox_inches="tight",\n'
              '      facecolor="white")\n'
              '    buf.seek(0)\n'
              '    slide.shapes.add_picture(\n'
              '        buf, l, t, w, h)\n'
              '    buf.close()\n'
              '\n'
              '# The FC matrix and HRF on the\n'
              '# left were generated live by\n'
              '# this script \u2014 no file on disk.',
              sz=12)

    _tb(s, Inches(0.6), Inches(5.5), Inches(7.5), Inches(0.8),
        "\u2191  These figures were rendered in-memory and embedded directly "
        "into this slide", sz=14, color=CLR_MED_GRAY)

    _notes(s, "The _mpl_to_slide helper renders to BytesIO and passes "
              "the buffer directly to add_picture. No temp file needed. "
              "Use dpi=200+ for crisp slides.")


def slide_figures_dpi(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, CLR_LIGHT_BG)
    _title_bar(s, "Figures: DPI and Resolution Guide")
    _slide_num(s)

    dpi_data = [
        ["Image Width on Slide", "Pixel Width @ 200 DPI", "@ 300 DPI"],
        ["Full width (11 in)", "2200 px", "3300 px"],
        ["Half width (5.5 in)", "1100 px", "1650 px"],
        ["Quarter (2.75 in)", "550 px", "825 px"],
    ]
    _add_table(s, Inches(0.8), Inches(1.3), Inches(6.5), Inches(2.5),
               dpi_data, header_color=CLR_TEAL, font_size=16)

    _rich(s, Inches(7.8), Inches(1.3), Inches(5.0), Inches(5.0), [
        ("Best practices:", True, 22),
        ("", False, 8),
        ("  \u2022  Use PNG for scientific figures (lossless)", False, 18),
        ("  \u2022  Use dpi=200 minimum from matplotlib", False, 18),
        ("  \u2022  Set facecolor='white' for light slides", False, 18),
        ("  \u2022  Set transparent=True for dark slides", False, 18),
        ("  \u2022  bbox_inches='tight' to crop whitespace", False, 18),
        ("", False, 8),
        ("Freeview screenshots:", True, 22),
        ("", False, 8),
        ("  1. Position the view in Freeview", False, 18),
        ("  2. File \u2192 Save Screenshot (.png)", False, 18),
        ("  3. Save to figures/capstone/", False, 18),
        ("  4. Reference in build_pptx.py", False, 18),
    ])

    _notes(s, "At 200 DPI, a full-width figure needs 2200px. PNG is "
              "always preferred over JPEG for scientific figures.")


# ── MATHEMATICS ───────────────────────────────────────────────────────────

def slide_math_intro(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, CLR_DARK)
    _tb(s, Inches(0.9), Inches(1.5), Inches(11), Inches(1.2),
        "Section 3", sz=28, color=CLR_ACCENT, font="Calibri Light")
    _tb(s, Inches(0.9), Inches(2.5), Inches(11), Inches(2),
        "Mathematics in Slides\nLaTeX-Style Typesetting",
        sz=48, bold=True, color=CLR_WHITE, font="Calibri Light")
    _tb(s, Inches(0.9), Inches(4.8), Inches(11), Inches(1.2),
        "Three approaches:  Unicode  \u00b7  LaTeX\u2192PNG  \u00b7  OOXML",
        sz=22, color=CLR_MED_GRAY, font="Calibri Light")


def slide_math_unicode(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, CLR_LIGHT_BG)
    _title_bar(s, "Math Approach A: Unicode Symbols")
    _slide_num(s)

    # Left: symbol reference table
    sym_data = [
        ["Symbol", "Unicode", "Python"],
        ["\u03b1  \u03b2  \u03b3  \u03b4", "Greek lowercase",
         "\\u03b1 \\u03b2 ..."],
        ["\u0394  \u03a9  \u03a3  \u03a0", "Greek uppercase",
         "\\u0394 \\u03a9 ..."],
        ["\u2207  \u2202  \u222b  \u2211", "Operators",
         "\\u2207 \\u2202 ..."],
        ["\u2192  \u21d2  \u2248  \u2260", "Relations / arrows",
         "\\u2192 \\u21d2 ..."],
        ["\u2080 \u2081 \u2082  \u00b2 \u00b3", "Sub / superscripts",
         "\\u2080 \\u00b2 ..."],
    ]
    _add_table(s, Inches(0.8), Inches(1.3), Inches(6.0), Inches(3.0),
               sym_data, header_color=CLR_TEAL, font_size=15,
               col_widths=[Inches(2.0), Inches(2.0), Inches(2.0)])

    # Right: live example
    _tb(s, Inches(7.3), Inches(1.3), Inches(5.5), Inches(0.5),
        "Live example (Cambria Math):", sz=18, bold=True)

    box = s.shapes.add_textbox(
        Inches(7.3), Inches(2.0), Inches(5.5), Inches(1.5))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "\u2207 \u00b7 \u03c3(u, \u03b8) + b = \u03c1 \u00fc"
    r.font.name = "Cambria Math"
    r.font.size = Pt(36)
    r.font.color.rgb = CLR_DARK_TEXT

    p2 = tf.add_paragraph()
    p2.space_before = Pt(12)
    r2 = p2.add_run()
    r2.text = "dz\u1d62/dt = (a\u1d62 + i\u03c9\u1d62)z\u1d62 \u2013 " \
              "|z\u1d62|\u00b2z\u1d62 + \u03b2\u2211SC\u1d62\u2c7c" \
              "(z\u2c7c \u2013 z\u1d62)"
    r2.font.name = "Cambria Math"
    r2.font.size = Pt(24)
    r2.font.color.rgb = CLR_DARK_TEXT

    _rich(s, Inches(0.8), Inches(4.6), Inches(11.5), Inches(2.5), [
        ("Pros:  ", True, 18, CLR_WM),
        ("  \u2022  Works in PowerPoint, Keynote, and Google Slides  "
         "  \u2022  No rendering step  "
         "  \u2022  Fully editable text", False, 16),
        ("Cons:  ", True, 18, CLR_TUMOR),
        ("  \u2022  No fractions, integrals with limits, or matrices  "
         "  \u2022  Limited sub/superscript set", False, 16),
        ("Tip:  ", True, 18, CLR_TEAL),
        ("  Use Cambria Math font for best Unicode math glyph coverage.",
         False, 16),
    ])

    _notes(s, "Unicode approach: simplest, most portable, but limited to "
              "symbols that have single codepoints. No stacking, fractions, "
              "or display-style integrals.")


def slide_math_latex(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, CLR_LIGHT_BG)
    _title_bar(s, "Math Approach B: LaTeX \u2192 PNG \u2192 Slide")
    _slide_num(s)

    _tb(s, Inches(0.8), Inches(1.2), Inches(6.5), Inches(0.5),
        "Rendered live by matplotlib \u2192 embedded as transparent PNG:",
        sz=18, bold=True)

    # Label + equation pairs
    _tb(s, Inches(0.8), Inches(1.8), Inches(3), Inches(0.4),
        "Momentum (weak form):", sz=16, color=CLR_TEAL, bold=True)
    _add_latex(s,
        r"\int_{\Omega} \mathbf{P} : \nabla \delta\mathbf{v} \,"
        r" d\Omega = \int_{\Omega} \mathbf{b} \cdot \delta\mathbf{v} \,"
        r" d\Omega",
        Inches(0.8), Inches(2.2), Inches(7), fontsize=28)

    _tb(s, Inches(0.8), Inches(3.2), Inches(4), Inches(0.4),
        "Strain energy (Neo-Hookean + Holzapfel):", sz=16,
        color=CLR_TEAL, bold=True)
    _add_latex(s,
        r"\Psi = \frac{\mu}{2}(I_1 - 3) + \frac{k_1}{2k_2}"
        r"\left(e^{k_2(I_4 - 1)^2} - 1\right)",
        Inches(0.8), Inches(3.6), Inches(6.5), fontsize=28)

    _tb(s, Inches(0.8), Inches(4.7), Inches(3), Inches(0.4),
        "Hopf bifurcation model:", sz=16, color=CLR_TEAL, bold=True)
    _add_latex(s,
        r"\frac{dz_i}{dt} = (a_i + i\omega_i)z_i"
        r" - |z_i|^2 z_i + \beta \sum_j SC_{ij}(z_j - z_i)"
        r" + \sigma\xi_i",
        Inches(0.8), Inches(5.1), Inches(7.5), fontsize=24)

    _code_box(s, Inches(8.0), Inches(1.3), Inches(4.8), Inches(5.3),
              'def _latex_to_buf(latex, fontsize,\n'
              '                  dpi=300, color="black"):\n'
              '    fig, ax = plt.subplots(\n'
              '        figsize=(0.01, 0.01))\n'
              '    ax.axis("off")\n'
              '    fig.patch.set_alpha(0.0)\n'
              '    txt = ax.text(0, 0, f"${latex}$",\n'
              '        fontsize=fontsize,\n'
              '        color=color, ...)\n'
              '    fig.canvas.draw()\n'
              '    bb = txt.get_window_extent(...)\n'
              '    fig.savefig(buf, format="png",\n'
              '        dpi=dpi, transparent=True,\n'
              '        bbox_inches=bb_inches)\n'
              '    return buf\n'
              '\n'
              '# Usage:\n'
              '_add_latex(slide,\n'
              '  r"\\Psi = \\frac{\\mu}{2}...",\n'
              '  Inches(1), Inches(3), Inches(6))\n'
              '\n'
              '# Note: \\begin{bmatrix} needs\n'
              '# plt.rcParams["text.usetex"]=True\n'
              '# (requires local LaTeX install)',
              sz=12)

    _notes(s, "These three equations are rendered live by matplotlib's "
              "mathtext engine and embedded as transparent PNGs. For full "
              "LaTeX support (\\bm, \\mathbb, AMS envs), set "
              "plt.rcParams['text.usetex'] = True (needs LaTeX installed).")


def slide_math_latex_dark(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, CLR_DARK)
    _tb(s, Inches(0), Inches(0), SLIDE_W, Inches(0.9),
        "  Math on Dark Backgrounds", sz=26, bold=True, color=CLR_WHITE,
        font="Calibri Light")
    s.shapes[-1].fill.solid()
    s.shapes[-1].fill.fore_color.rgb = RGBColor(0x00, 0x55, 0x66)
    _slide_num(s)

    _tb(s, Inches(0.8), Inches(1.2), Inches(5), Inches(0.4),
        "Pass color='white' to _latex_to_buf:", sz=18, color=CLR_MED_GRAY)

    _add_latex_dark(s,
        r"\nabla \cdot \sigma(\mathbf{u}, \theta)"
        r" + \mathbf{b} = \rho \ddot{\mathbf{u}}",
        Inches(0.8), Inches(1.8), Inches(8), fontsize=34)

    _add_latex_dark(s,
        r"\frac{\partial \mathbf{u}}{\partial t} = "
        r"\mathbf{f}(\mathbf{u}, SC, \theta)",
        Inches(0.8), Inches(3.0), Inches(6), fontsize=30)

    _add_latex_dark(s,
        r"SC' = SC \cdot \exp("
        r"-\alpha\,|\mathbf{u}|)",
        Inches(0.8), Inches(4.2), Inches(7), fontsize=30)

    _add_latex_dark(s,
        r"a_i' = a_i - \gamma \cdot \epsilon_i",
        Inches(0.8), Inches(5.4), Inches(5), fontsize=30)

    _tb(s, Inches(0.8), Inches(6.2), Inches(11), Inches(0.6),
        "Transparent background ensures equations float over any slide "
        "colour.", sz=16, color=CLR_MED_GRAY, font="Calibri Light")

    _notes(s, "For dark backgrounds, set color='white' in the LaTeX "
              "rendering call. transparent=True in savefig ensures no "
              "white rectangle behind the equation.")


def slide_math_ooxml(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, CLR_LIGHT_BG)
    _title_bar(s, "Math Approach C: Native OOXML Equations")
    _slide_num(s)

    _rich(s, Inches(0.8), Inches(1.3), Inches(5.5), Inches(2.5), [
        ("Editable equations inside PowerPoint:", True, 20),
        ("", False, 8),
        ("  \u2022  Insert OOXML <m:oMath> elements via lxml", False, 18),
        ("  \u2022  Right-click \u2192 Edit Equation in PowerPoint", False, 18),
        ("  \u2022  Most verbose approach", False, 18),
        ("  \u2022  Best for templates others will edit", False, 18),
    ])

    _code_box(s, Inches(0.8), Inches(3.5), Inches(5.5), Inches(3.2),
              '<m:oMath xmlns:m="...officeDocument\n'
              '         /2006/math">\n'
              '  <m:r><m:t>\u2207</m:t></m:r>\n'
              '  <m:r><m:t> \u00b7 </m:t></m:r>\n'
              '  <m:r>\n'
              '    <m:rPr>\n'
              '      <m:sty m:val="bi"/>\n'
              '    </m:rPr>\n'
              '    <m:t>\u03c3</m:t>\n'
              '  </m:r>\n'
              '</m:oMath>', sz=13)

    # Comparison table
    cmp = [
        ["", "Unicode (A)", "LaTeX\u2192PNG (B)", "OOXML (C)"],
        ["Complexity", "Trivial", "Moderate", "High"],
        ["Editable?", "Yes (text)", "No (image)", "Yes (equation)"],
        ["Fractions", "No", "Yes", "Yes"],
        ["Matrices", "No", "Yes", "Yes"],
        ["Portability", "Everywhere", "Everywhere", "PowerPoint only"],
        ["Recommended", "Simple expr.", "Complex math", "Shared templates"],
    ]
    _add_table(s, Inches(6.8), Inches(1.3), Inches(6.0), Inches(4.5),
               cmp, header_color=CLR_TEAL, font_size=14,
               col_widths=[Inches(1.3), Inches(1.5), Inches(1.7),
                           Inches(1.5)])

    _notes(s, "OOXML math is the most complex but produces editable "
              "equations. For most scientific talks, LaTeX→PNG (Approach B) "
              "is the sweet spot.")


# ── MIXED FORMATTING & HYPERLINKS ─────────────────────────────────────────

def slide_formatting_intro(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, CLR_DARK)
    _tb(s, Inches(0.9), Inches(1.5), Inches(11), Inches(1.2),
        "Section 4", sz=28, color=CLR_ACCENT, font="Calibri Light")
    _tb(s, Inches(0.9), Inches(2.5), Inches(11), Inches(2),
        "Mixed Formatting,\nHyperlinks, and Slide Numbers",
        sz=48, bold=True, color=CLR_WHITE, font="Calibri Light")
    _tb(s, Inches(0.9), Inches(4.8), Inches(11), Inches(1.2),
        "Multiple runs per paragraph  \u00b7  clickable links  "
        "\u00b7  field codes",
        sz=20, color=CLR_MED_GRAY, font="Calibri Light")


def slide_mixed_demo(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, CLR_LIGHT_BG)
    _title_bar(s, "Mixed Formatting: Multiple Runs per Paragraph")
    _slide_num(s)

    # Demo 1: bold + colour + size in one line
    _tb(s, Inches(0.8), Inches(1.2), Inches(5), Inches(0.4),
        "Example 1 \u2014 Inline emphasis:", sz=18, bold=True, color=CLR_TEAL)

    box1 = s.shapes.add_textbox(
        Inches(0.8), Inches(1.7), Inches(8), Inches(0.6))
    p1 = box1.text_frame.paragraphs[0]
    _mixed_run(p1, "Tumor volume: ", sz=22)
    _mixed_run(p1, "50.2 cm\u00b3", sz=22, bold=True, color=CLR_TUMOR)
    _mixed_run(p1, "  (manual segmentation)", sz=16, color=CLR_MED_GRAY,
               italic=True)

    # Demo 2: code-style inline
    _tb(s, Inches(0.8), Inches(2.5), Inches(5), Inches(0.4),
        "Example 2 \u2014 Inline code:", sz=18, bold=True, color=CLR_TEAL)

    box2 = s.shapes.add_textbox(
        Inches(0.8), Inches(3.0), Inches(8), Inches(0.6))
    p2 = box2.text_frame.paragraphs[0]
    _mixed_run(p2, "Run ", sz=20)
    _mixed_run(p2, "recon-all-clinical", sz=20, font="Menlo",
               color=CLR_TEAL, bold=True)
    _mixed_run(p2, " for post-op segmentation", sz=20)

    # Demo 3: hyperlinks
    _tb(s, Inches(0.8), Inches(3.8), Inches(5), Inches(0.4),
        "Example 3 \u2014 Clickable hyperlinks:", sz=18, bold=True,
        color=CLR_TEAL)

    box3 = s.shapes.add_textbox(
        Inches(0.8), Inches(4.3), Inches(10), Inches(0.6))
    p3 = box3.text_frame.paragraphs[0]
    _mixed_run(p3, "Data available at ", sz=20)
    _mixed_run(p3, "OpenNeuro ds002080", sz=20, color=CLR_CAVITY,
               underline=True,
               hyperlink="https://openneuro.org/datasets/ds002080")
    _mixed_run(p3, " and ", sz=20)
    _mixed_run(p3, "GitHub", sz=20, color=CLR_CAVITY, underline=True,
               hyperlink="https://github.com/arvidl/glioma-brain-digital-twin-public")

    # Code on the right
    _code_box(s, Inches(7.5), Inches(1.2), Inches(5.4), Inches(5.0),
              '# Multiple runs in one paragraph\n'
              'p = tf.paragraphs[0]\n'
              '\n'
              'r1 = p.add_run()\n'
              'r1.text = "Tumor volume: "\n'
              'r1.font.size = Pt(22)\n'
              '\n'
              'r2 = p.add_run()\n'
              'r2.text = "50.2 cm\u00b3"\n'
              'r2.font.bold = True\n'
              'r2.font.color.rgb = CLR_TUMOR\n'
              '\n'
              '# Hyperlink\n'
              'r = p.add_run()\n'
              'r.text = "OpenNeuro ds002080"\n'
              'r.font.underline = True\n'
              'r.hyperlink.address = \\\n'
              '  "https://openneuro.org/..."',
              sz=12)

    _notes(s, "Each paragraph can contain multiple runs with independent "
              "formatting. This is how you get bold + colour + size "
              "variations on a single line. Hyperlinks are set via "
              "run.hyperlink.address.")


def slide_slide_numbers(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, CLR_LIGHT_BG)
    _title_bar(s, "Slide Numbering via OOXML Field Codes")
    _slide_num(s)

    _rich(s, Inches(0.8), Inches(1.3), Inches(5.5), Inches(2.5), [
        ("Automatic slide numbers:", True, 22),
        ("", False, 8),
        ("  \u2022  Uses OOXML <a:fld type=\"slidenum\">", False, 18),
        ("  \u2022  Updates automatically when slides reorder", False, 18),
        ("  \u2022  Placed in bottom-right on every content slide", False, 18),
        ("  \u2022  Skipped on section dividers (design choice)", False, 18),
    ])

    _code_box(s, Inches(0.8), Inches(3.5), Inches(5.5), Inches(3.2),
              'from pptx.oxml.ns import qn\n'
              'from lxml import etree\n'
              '\n'
              'def _slide_num(slide):\n'
              '    box = slide.shapes.add_textbox(\n'
              '      Inches(12.4), Inches(7),\n'
              '      Inches(0.8), Inches(0.4))\n'
              '    p = box.text_frame.paragraphs[0]\n'
              '    p.alignment = PP_ALIGN.RIGHT\n'
              '    fld = etree.SubElement(\n'
              '      p._p, qn("a:fld"),\n'
              '      attrib={"type": "slidenum",\n'
              '              "id": "{UUID}"})\n'
              '    rPr = etree.SubElement(\n'
              '      fld, qn("a:rPr"), ...)\n'
              '    t = etree.SubElement(\n'
              '      fld, qn("a:t"))\n'
              '    t.text = "<#>"',
              sz=12)

    _tb(s, Inches(6.8), Inches(1.3), Inches(6), Inches(3),
        "Look at the bottom-right corner \u2192\n\n"
        "Every content slide in this tutorial has an auto-updating "
        "slide number placed by the _slide_num() helper.\n\n"
        "The number is styled in grey to stay unobtrusive.",
        sz=18, color=CLR_DARK_TEXT)

    _notes(s, "Field codes are OOXML elements that PowerPoint evaluates "
              "at render time. slidenum is the most common, but there are "
              "also datetime and slidetotal field types.")


# ── WORKFLOW ──────────────────────────────────────────────────────────────

def slide_workflow(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, CLR_LIGHT_BG)
    _title_bar(s, "Workflow: Edit \u2192 Generate \u2192 Present \u2192 Iterate")
    _slide_num(s)

    _code_box(s, Inches(0.8), Inches(1.3), Inches(7.5), Inches(5.0),
              "\u250c\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500"
              "\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500"
              "\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500"
              "\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500"
              "\u2500\u2500\u2500\u2510\n"
              "\u2502 1. Run capstone notebooks (NB_C01\u2013C09)    \u2502\n"
              "\u2502    \u2192 figures saved to                     \u2502\n"
              "\u2502      notebooks/figures/capstone/             \u2502\n"
              "\u2502                                              \u2502\n"
              "\u2502 2. Capture Freeview screenshots               \u2502\n"
              "\u2502    \u2192 save to same figures directory       \u2502\n"
              "\u2502                                              \u2502\n"
              "\u2502 3. Edit build_pptx.py                        \u2502\n"
              "\u2502    \u2022 Add/remove slides in build_cNN()    \u2502\n"
              "\u2502    \u2022 Reference new figures by filename   \u2502\n"
              "\u2502    \u2022 Write speaker notes                 \u2502\n"
              "\u2502                                              \u2502\n"
              "\u2502 4. python beamer/build_pptx.py               \u2502\n"
              "\u2502                                              \u2502\n"
              "\u2502 5. Open in PowerPoint / Keynote / Slides     \u2502\n"
              "\u2502    \u2022 Fine-tune alignment & animations    \u2502\n"
              "\u2502    \u2022 Add live demo or video links        \u2502\n"
              "\u2502                                              \u2502\n"
              "\u2502 6. Iterate: go back to step 3                \u2502\n"
              "\u2514\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500"
              "\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500"
              "\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500"
              "\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500\u2500"
              "\u2500\u2500\u2500\u2518",
              sz=14)

    _rich(s, Inches(8.8), Inches(1.3), Inches(4.0), Inches(5.5), [
        ("Why script your slides?", True, 22, CLR_TEAL),
        ("", False, 8),
        ("  \u2022  Reproducible from source", False, 18),
        ("  \u2022  Version-controlled (git diff)", False, 18),
        ("  \u2022  Figures auto-update", False, 18),
        ("  \u2022  Consistent formatting", False, 18),
        ("  \u2022  Speaker notes in code", False, 18),
        ("  \u2022  CI/CD: auto-build on push", False, 18),
        ("", False, 8),
        ("Complement with:", True, 22, CLR_TEAL),
        ("", False, 8),
        ("  \u2022  Beamer for LaTeX-native math", False, 18),
        ("  \u2022  Google Slides for Freeview", False, 18),
        ("    screenshots and collaboration", False, 18),
    ])

    _notes(s, "The key advantage of scripted slides: regenerate the entire "
              "deck after updating a figure or fixing a typo, without "
              "touching PowerPoint manually.")


# ── REFERENCES ────────────────────────────────────────────────────────────

def slide_references(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, CLR_DARK)

    _tb(s, Inches(0.9), Inches(0.8), Inches(11), Inches(1),
        "References and Resources", sz=42, bold=True, color=CLR_WHITE,
        font="Calibri Light")

    box = s.shapes.add_textbox(
        Inches(0.9), Inches(2.0), Inches(11), Inches(4.5))
    tf = box.text_frame
    tf.word_wrap = True

    links = [
        ("python-pptx documentation",
         "https://python-pptx.readthedocs.io/"),
        ("API reference \u2014 Shapes",
         "https://python-pptx.readthedocs.io/en/latest/api/shapes.html"),
        ("API reference \u2014 Tables",
         "https://python-pptx.readthedocs.io/en/latest/api/table.html"),
        ("Working with text",
         "https://python-pptx.readthedocs.io/en/latest/user/text.html"),
        ("OOXML math spec (ECMA-376)",
         "https://ecma-international.org/publications-and-standards/"
         "standards/ecma-376/"),
        ("Unicode mathematical symbols",
         "https://en.wikipedia.org/wiki/Mathematical_operators_and_"
         "symbols_in_Unicode"),
        ("matplotlib mathtext tutorial",
         "https://matplotlib.org/stable/tutorials/text/mathtext.html"),
        ("This project: glioma-brain-digital-twin",
         "https://github.com/arvidl/glioma-brain-digital-twin-public"),
    ]

    for i, (label, url) in enumerate(links):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = f"  \u2022  {label}"
        r.font.size = Pt(18)
        r.font.name = "Calibri Light"
        r.font.color.rgb = CLR_CAVITY
        r.font.underline = True
        r.hyperlink.address = url

    _notes(s, "All links are clickable in presentation mode. The python-pptx "
              "docs are the authoritative reference; the ECMA-376 spec is "
              "needed for OOXML math.")


# ═══════════════════════════════════════════════════════════════════════════
# Main
# ═══════════════════════════════════════════════════════════════════════════

def build_tutorial(output: Path):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slide_title(prs)
    slide_overview(prs)
    slide_architecture(prs)

    slide_tables_intro(prs)
    slide_tables_demo(prs)
    slide_tables_advanced(prs)

    slide_figures_intro(prs)
    slide_figures_mpl(prs)
    slide_figures_dpi(prs)

    slide_math_intro(prs)
    slide_math_unicode(prs)
    slide_math_latex(prs)
    slide_math_latex_dark(prs)
    slide_math_ooxml(prs)

    slide_formatting_intro(prs)
    slide_mixed_demo(prs)
    slide_slide_numbers(prs)

    slide_workflow(prs)
    slide_references(prs)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    print(f"Saved {len(prs.slides)} slides \u2192 {output}")


def main():
    ap = argparse.ArgumentParser(
        description="Build the python-pptx tutorial slide deck.")
    ap.add_argument("-o", "--output", type=Path, default=DEFAULT_OUTPUT,
                    help="Output path (default: %(default)s)")
    args = ap.parse_args()
    build_tutorial(args.output)


if __name__ == "__main__":
    main()
