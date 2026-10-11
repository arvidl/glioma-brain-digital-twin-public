#!/usr/bin/env python3
"""
build_pptx.py  --  Programmatically build a PowerPoint slide deck for the
Glioma Brain Digital Twin presentation (Finse, March 2026).

Usage
-----
    python beamer/build_pptx.py                   # default output
    python beamer/build_pptx.py -o talk.pptx      # custom output path
    python beamer/build_pptx.py --fig-dir /path    # alternate figure root

The script mirrors the structure of the nine capstone notebooks (NB_C01–C09)
and the existing Beamer slide deck, producing a self-contained .pptx file
with embedded figures, speaker notes, and a consistent visual identity.

Requires: python-pptx  (pip install python-pptx)
"""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_FIG_DIR = REPO_ROOT / "notebooks" / "figures" / "capstone"
DEFAULT_OUTPUT = REPO_ROOT / "beamer" / "glioma_digital_twin_finse2026.pptx"
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# Project colours (from beamer/preamble.tex)
CLR_DARK = RGBColor(0x23, 0x37, 0x3B)      # dark teal (backgrounds)
CLR_TEAL = RGBColor(0x00, 0x77, 0x88)      # metropolis teal
CLR_ACCENT = RGBColor(0xEB, 0x81, 0x1E)    # warm orange accent
CLR_TUMOR = RGBColor(0xCC, 0x33, 0x33)     # tumor red
CLR_CAVITY = RGBColor(0x33, 0x66, 0xCC)    # cavity blue
CLR_WM = RGBColor(0x33, 0x99, 0x66)        # white-matter green
CLR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CLR_LIGHT_BG = RGBColor(0xF5, 0xF5, 0xF0)  # off-white for content slides
CLR_MED_GRAY = RGBColor(0x88, 0x88, 0x88)
CLR_DARK_TEXT = RGBColor(0x2D, 0x2D, 0x2D)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _set_slide_bg(slide, color: RGBColor):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_textbox(slide, left, top, width, height, text, *,
                 font_size=18, bold=False, color=CLR_DARK_TEXT,
                 alignment=PP_ALIGN.LEFT, font_name="Calibri",
                 anchor=MSO_ANCHOR.TOP):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf.paragraphs[0].alignment = alignment
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    try:
        txBox.text_frame._txBody.attrib[
            "{http://schemas.openxmlformats.org/drawingml/2006/main}anchor"
        ] = "t"
    except Exception:
        pass
    return txBox


def _add_rich_textbox(slide, left, top, width, height, paragraphs, *,
                      font_size=18, color=CLR_DARK_TEXT, font_name="Calibri",
                      line_spacing=1.3, alignment=PP_ALIGN.LEFT):
    """paragraphs: list of (text, bold, font_size_override|None)."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, (txt, bld, sz_override) in enumerate(paragraphs):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = alignment
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = txt
        run.font.size = Pt(sz_override or font_size)
        run.font.bold = bld
        run.font.color.rgb = color
        run.font.name = font_name
    return txBox


def _add_image(slide, fig_dir: Path, filename: str,
               left, top, width, height=None):
    """Insert an image if the file exists; otherwise add a placeholder box."""
    img_path = fig_dir / filename
    if img_path.exists():
        if height:
            slide.shapes.add_picture(str(img_path), left, top, width, height)
        else:
            slide.shapes.add_picture(str(img_path), left, top, width=width)
    else:
        _add_textbox(
            slide, left, top, width, height or Inches(3),
            f"[Figure: {filename}]",
            font_size=14, color=CLR_MED_GRAY,
            alignment=PP_ALIGN.CENTER,
        )


def _set_notes(slide, text: str):
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = text


def _section_divider(prs, number: str, title: str, subtitle: str,
                     notes: str = ""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    _set_slide_bg(slide, CLR_DARK)
    _add_textbox(slide, Inches(0.8), Inches(1.5), Inches(11), Inches(1.2),
                 number, font_size=28, bold=False, color=CLR_ACCENT,
                 font_name="Calibri Light")
    _add_textbox(slide, Inches(0.8), Inches(2.4), Inches(11), Inches(2),
                 title, font_size=44, bold=True, color=CLR_WHITE,
                 font_name="Calibri Light")
    _add_textbox(slide, Inches(0.8), Inches(4.6), Inches(11), Inches(1.5),
                 subtitle, font_size=22, bold=False, color=CLR_MED_GRAY,
                 font_name="Calibri Light")
    if notes:
        _set_notes(slide, notes)
    return slide


def _content_slide(prs, title: str, bullets: list[str], fig_dir: Path,
                   figure: str | None = None, notes: str = "",
                   two_col: bool = True):
    """Content slide: title bar + bullets (left) + optional figure (right)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    _set_slide_bg(slide, CLR_LIGHT_BG)

    # Title bar
    _add_textbox(slide, Inches(0), Inches(0), SLIDE_WIDTH, Inches(0.9),
                 f"  {title}", font_size=26, bold=True, color=CLR_WHITE,
                 font_name="Calibri Light")
    title_bg = slide.shapes[len(slide.shapes) - 1]
    title_bg.fill.solid()
    title_bg.fill.fore_color.rgb = CLR_TEAL

    # Bullet column
    if two_col and figure:
        bul_width = Inches(5.8)
    else:
        bul_width = Inches(11.5)

    bullet_pars = []
    for b in bullets:
        bullet_pars.append((f"  \u2022  {b}", False, 18))
    _add_rich_textbox(
        slide, Inches(0.6), Inches(1.2), bul_width, Inches(5.5),
        bullet_pars, color=CLR_DARK_TEXT,
    )

    # Figure (right column)
    if figure:
        fig_left = Inches(6.8) if two_col else Inches(1.5)
        fig_top = Inches(1.3) if two_col else Inches(4.0)
        fig_w = Inches(5.8) if two_col else Inches(10)
        _add_image(slide, fig_dir, figure, fig_left, fig_top, fig_w,
                   Inches(5.2) if two_col else None)

    if notes:
        _set_notes(slide, notes)
    return slide


def _full_figure_slide(prs, title: str, fig_dir: Path, figure: str,
                       caption: str = "", notes: str = ""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide, CLR_LIGHT_BG)

    _add_textbox(slide, Inches(0), Inches(0), SLIDE_WIDTH, Inches(0.9),
                 f"  {title}", font_size=26, bold=True, color=CLR_WHITE,
                 font_name="Calibri Light")
    title_bg = slide.shapes[len(slide.shapes) - 1]
    title_bg.fill.solid()
    title_bg.fill.fore_color.rgb = CLR_TEAL

    _add_image(slide, fig_dir, figure,
               Inches(1.0), Inches(1.2), Inches(11.3), Inches(5.3))

    if caption:
        _add_textbox(slide, Inches(1.0), Inches(6.6), Inches(11.3),
                     Inches(0.6), caption, font_size=14,
                     color=CLR_MED_GRAY, alignment=PP_ALIGN.CENTER)
    if notes:
        _set_notes(slide, notes)
    return slide


# ---------------------------------------------------------------------------
# Slide builders per section
# ---------------------------------------------------------------------------

def build_title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide, CLR_DARK)

    _add_textbox(slide, Inches(0.8), Inches(1.4), Inches(11), Inches(1.8),
                 "Glioma Brain Digital Twin",
                 font_size=52, bold=True, color=CLR_WHITE,
                 font_name="Calibri Light")
    _add_textbox(slide, Inches(0.8), Inches(3.4), Inches(11), Inches(1.2),
                 "From Multiparametric MRI to Mathematical Challenges\n"
                 "in Brain Mechanics",
                 font_size=26, bold=False, color=CLR_ACCENT,
                 font_name="Calibri Light")
    _add_textbox(slide, Inches(0.8), Inches(5.2), Inches(11), Inches(1.5),
                 "Arvid Lundervold\n"
                 "CAS \u2014 Mathematical Challenges in Brain Mechanics\n"
                 "Finse, March 2026",
                 font_size=20, bold=False, color=CLR_MED_GRAY,
                 font_name="Calibri Light")
    _set_notes(slide,
               "Welcome and overview. This talk accompanies capstone "
               "notebooks NB_C01\u2013C09 from the glioma-brain-digital-twin "
               "project. All figures can be reproduced from those notebooks.")


def build_c01(prs, fig_dir):
    """NB_C01: Digital Twin Concept and Data Landscape."""
    _section_divider(
        prs, "Part I", "The Digital Twin Concept\nand Data Landscape",
        "NB_C01 \u2014 What is a digital twin?  The glioma case.  "
        "Data from multiple sources.",
        notes="NB_C01 covers the digital twin concept, mathematical "
              "challenges, and the multimodal data collected for PAT16.")

    _content_slide(prs,
        "What Is a Digital Twin?",
        ["Subject-specific computational replica",
         "Personalized geometry (T1w), fibers (dMRI), function (fMRI)",
         "Predicts response to perturbations (surgery, treatment)",
         "Validation against longitudinal observations",
         "Five layers: Geometric \u2192 Constitutive \u2192 Functional "
         "\u2192 Simulation \u2192 Clinical"],
        fig_dir, "figC01_01_coupling_chain.png",
        notes="The coupling chain: from patient data through five layers "
              "to clinical outcome prediction.")

    _content_slide(prs,
        "Why Gliomas?",
        ["Clinical urgency: GBM median survival ~15 months",
         "Mass effect, surgical brain shift (10\u201320 mm)",
         "Natural longitudinal experiment: pre-op \u2192 post-op",
         "Pathological anatomy breaks standard pipelines",
         "Rich multimodal data from the BTC-preop dataset (OpenNeuro)"],
        fig_dir, None,
        notes="Gliomas offer a unique opportunity: the disease and surgery "
              "create measurable deformations that serve as empirical "
              "validation data.")

    _content_slide(prs,
        "Mathematical Challenges",
        ["Ill-posed inverse problem: non-unique (\u03b8, BC)",
         "Constitutive modeling from dMRI (crossing fibers, infiltration)",
         "Registration error amplification in derivatives",
         "Boundary conditions at cavity\u2013brain interface",
         "Multimodal fusion across spatial and temporal scales",
         "Uncertainty quantification without MRE ground truth",
         "Pathological tissue: topology change, remodeling"],
        fig_dir, None,
        notes="Seven + one mathematical challenges identified for the "
              "glioma brain digital twin.")

    _full_figure_slide(prs,
        "Data Landscape: Patient PAT16",
        fig_dir, "figC01_02_data_landscape.png",
        "Three MRI modalities \u00d7 two timepoints \u2192 "
        "derivatives \u2192 digital twin layers",
        notes="Overview of data flowing from raw mpMRI through processing "
              "pipelines into the digital twin.")

    _content_slide(prs,
        "The Governing Equation",
        ["\u2207 \u00b7 \u03c3(u, \u03b8) + b = \u03c1 \u00fc",
         "MRI informs: domain \u03a9, parameters \u03b8, boundary conditions",
         "Forward problem: given (\u03b8, BC), predict u",
         "Inverse problem: given observed deformation, infer (\u03b8, BC)",
         "Non-uniqueness \u2192 need for regularization and priors"],
        fig_dir, None,
        notes="The momentum equation is central. MRI data inform every "
              "component: geometry, material parameters, and BCs.")


def build_c02(prs, fig_dir):
    """NB_C02: Clinical, Cognitive, and Psychosocial Profile."""
    _section_divider(
        prs, "Part II", "Clinical, Cognitive, and\nPsychosocial Profile",
        "NB_C02 \u2014 CANTAB, demographics, psychosocial measures",
        notes="NB_C02 covers the clinical and cognitive data that form "
              "Layer 5 (Clinical Outcome) of the digital twin.")

    _content_slide(prs,
        "The CANTAB Battery",
        ["Computerised cognitive tests: MOT, RTI, RVP, SOC, SSP",
         "Reaction time, sustained attention, spatial planning, "
         "working memory",
         "Each test maps to specific brain regions (DK68 parcellation)",
         "Pre-op vs post-op comparison \u2192 \u0394CANTAB"],
        fig_dir, "figC02_01_domain_parcel_mapping.png",
        notes="CANTAB tests are standardised, computerised, and have "
              "well-established neural correlates.")

    _full_figure_slide(prs,
        "Cognitive Change: Pre-op vs Post-op",
        fig_dir, "figC02_02_cantab_prepost.png",
        "Z-score changes across CANTAB subtests for PAT16",
        notes="Direction and magnitude of cognitive change across "
              "domains. SOC and SSP show the largest decrements.")

    _content_slide(prs,
        "Psychosocial Context",
        ["STAI (anxiety), PSWQ (worry), ERQ (emotion regulation), "
         "UCLA (loneliness)",
         "Elevated anxiety and worry can depress cognitive scores",
         "Post-surgical fatigue and stress are confounds",
         "Interpreting \u0394CANTAB requires psychosocial context"],
        fig_dir, "figC02_03_psychosocial.png",
        notes="Psychosocial measures are essential for interpreting "
              "cognitive outcomes; anxiety alone can shift z-scores.")


def build_c03(prs, fig_dir):
    """NB_C03: Structural MRI \u2014 Anatomy, Segmentation, Lesions."""
    _section_divider(
        prs, "Part III", "Structural MRI\nSegmentation and Lesions",
        "NB_C03 \u2014 T1w, FreeSurfer parcellation, tumor and cavity masks",
        notes="NB_C03 covers structural MRI: T1w visualization, "
              "segmentation, and lesion identification.")

    _full_figure_slide(prs,
        "T1w Pre-operative MRI",
        fig_dir, "figC03_01_t1w_preop.png",
        "Axial, coronal, sagittal views of pre-op T1w (PAT16)",
        notes="The pre-op T1w shows the tumor (right frontal) as a "
              "hypointense mass with surrounding edema.")

    _full_figure_slide(prs,
        "T1w Post-operative MRI",
        fig_dir, "figC03_02_t1w_postop.png",
        "Post-operative T1w showing the resection cavity",
        notes="The resection cavity replaces the tumor volume. "
              "Note the brain shift and ventricular remodeling.")

    _content_slide(prs,
        "Tumor Segmentation",
        ["Pre-op: binary tumor mask (manual + refined)",
         "Post-op: resection cavity via SynthSeg or manual",
         "Cavity masking critical for longitudinal registration",
         "Without masking: ANTs SyN collapses tissue into void",
         "Tumor volume: ~50 cm\u00b3 \u2192 cavity + remodeling"],
        fig_dir, "figC03_03_tumor_overlay.png",
        notes="Accurate lesion segmentation is the foundation for "
              "cavity-aware registration and FE mesh generation.")

    _full_figure_slide(prs,
        "FreeSurfer Parcellation (aparc+aseg)",
        fig_dir, "figC03_04_parcellation.png",
        "Desikan\u2013Killiany 68-region cortical + subcortical parcellation",
        notes="aparc+aseg provides 83 ROIs used for both structural "
              "and functional connectivity matrices.")

    _full_figure_slide(prs,
        "Pre-op vs Post-op Comparison",
        fig_dir, "figC03_05_preop_vs_postop.png",
        "Side-by-side: T1w, segmentation, and lesion masks",
        notes="Direct comparison of anatomy before and after surgery. "
              "Key for understanding brain shift and tissue remodeling.")


def build_c04(prs, fig_dir):
    """NB_C04: Diffusion MRI \u2014 Microstructure, Connectivity, Tractography."""
    _section_divider(
        prs, "Part IV", "Diffusion MRI\nMicrostructure and Tractography",
        "NB_C04 \u2014 DTI maps, structural connectivity, fiber tracking",
        notes="NB_C04 covers diffusion MRI principles, DTI scalar maps, "
              "structural connectivity, and tractography.")

    _content_slide(prs,
        "Diffusion MRI Principles",
        ["Water diffusion probes tissue microstructure at \u03bcm scale",
         "Multi-shell acquisition: b = 700, 1000, 2000 s/mm\u00b2",
         "AP + PA phase-encoding for susceptibility correction",
         "QSIPrep: topup, eddy, motion correction, quality control",
         "Tensor model \u2192 FA, MD, AD, RD; CSD \u2192 fODF \u2192 "
         "tractography"],
        fig_dir, None,
        notes="Multi-shell diffusion data enables both DTI analysis and "
              "constrained spherical deconvolution for fiber tracking.")

    _full_figure_slide(prs,
        "DTI Scalar Maps",
        fig_dir, "figC04_01_dti_maps.png",
        "FA, MD, AD, RD maps in T1w space (pre-op)",
        notes="DTI maps co-registered to T1w space. FA shows anisotropy "
              "(white matter bright), MD shows mean diffusivity.")

    _full_figure_slide(prs,
        "DTI Value Distributions",
        fig_dir, "figC04_02_dti_histograms.png",
        "Histograms of FA and MD by tissue class (GM, WM, tumor)",
        notes="Distribution differences between tissue classes inform "
              "the constitutive model: WM is stiffer and more anisotropic.")

    _content_slide(prs,
        "Structural Connectivity",
        ["MRtrix3: CSD \u2192 ACT \u2192 SIFT2 \u2192 83\u00d783 connectome",
         "SC matrix: weighted streamline counts between DK68 parcels",
         "SIFT2 improves biological plausibility of weights",
         "Pre-op vs post-op: tract disruption near resection",
         "SC informs the neural dynamics model (NB_C07)"],
        fig_dir, "figC04_03_structural_connectivity.png",
        notes="The structural connectivity matrix is the backbone of "
              "the neural dynamics simulation.")

    _content_slide(prs,
        "Material Properties from dMRI",
        ["DTI \u2192 fiber direction a\u2080 and anisotropy \u03ba",
         "Neo-Hookean + Holzapfel constitutive model",
         "Region-wise: GM (\u03bc=1 kPa), WM (\u03bc=1.5 kPa), "
         "deep nuclei, CSF",
         "Tumor zone: altered stiffness (infiltration uncertainty)",
         "Crossing fibers: CSD provides multiple directions"],
        fig_dir, "figC04_04_material_properties.png",
        notes="dMRI-derived material properties feed Layer 2 of the "
              "digital twin.")


def build_c05(prs, fig_dir):
    """NB_C05: Resting-State fMRI \u2014 HRF, rsHRF, Functional Connectivity."""
    _section_divider(
        prs, "Part V", "Resting-State fMRI\nHemodynamics and Connectivity",
        "NB_C05 \u2014 BOLD signal, rsHRF deconvolution, FC matrices",
        notes="NB_C05 covers rs-fMRI: BOLD signal, HRF, rsHRF, and "
              "functional connectivity.")

    _content_slide(prs,
        "The BOLD Signal and Neurovascular Coupling",
        ["BOLD: blood-oxygen-level dependent contrast",
         "Indirect measure of neural activity via hemodynamic response",
         "Canonical HRF: ~6 s peak, ~16 s undershoot",
         "Resting-state HRF (rsHRF): voxel-wise estimation",
         "Point-process deconvolution of spontaneous events",
         "fMRIPrep: confound regression, bandpass filtering"],
        fig_dir, None,
        notes="The BOLD signal is an indirect, delayed, and spatially "
              "blurred measure of neural activity.")

    _full_figure_slide(prs,
        "rsHRF Parameter Maps",
        fig_dir, "figC05_01_rshrf_maps.png",
        "Voxel-wise rsHRF height, time-to-peak, and FWHM",
        notes="rsHRF maps reveal regional hemodynamic variability, "
              "potentially altered near the tumor.")

    _content_slide(prs,
        "Functional Connectivity",
        ["Pearson correlation between ROI time series \u2192 83\u00d783 FC",
         "Fisher-z transform for statistical comparison",
         "Pre-op vs post-op: reorganisation of functional networks",
         "FC degree: hub regions and their disruption",
         "SC\u2013FC coupling: structural basis of functional connections"],
        fig_dir, "figC05_02_functional_connectivity.png",
        notes="The FC matrix captures functional relationships between "
              "brain regions at rest.")

    _full_figure_slide(prs,
        "Functional Connectivity Degree",
        fig_dir, "figC05_03_fc_degree.png",
        "Node degree changes pre-op \u2192 post-op",
        notes="FC degree highlights hubs and reveals post-surgical "
              "functional reorganization.")


def build_c06(prs, fig_dir):
    """NB_C06: Brain Mechanics \u2014 Mass Effect and Tissue Deformation."""
    _section_divider(
        prs, "Part VI", "Brain Mechanics\nMass Effect and Deformation",
        "NB_C06 \u2014 Continuum mechanics, FEM, displacement fields",
        notes="NB_C06 covers the brain mechanics forward problem: "
              "constitutive model, FE mesh, simulation, and validation.")

    _content_slide(prs,
        "Continuum Mechanics of Brain Tissue",
        ["Hyperelastic: Neo-Hookean isotropic + Holzapfel anisotropic",
         "Fiber direction from DTI eigenvectors",
         "Large deformation (finite strain): F, C, J",
         "Piola\u2013Kirchhoff stress \u2192 weak form",
         "Newton solver via PETSc/SNES in FEniCSx",
         "BCs: gravity body force, skull (fixed), ventricles, "
         "tumor/cavity interface"],
        fig_dir, None,
        notes="The constitutive model combines isotropic ground matrix "
              "with anisotropic fiber reinforcement from DTI.")

    _content_slide(prs,
        "Mass Effect and Brain Shift",
        ["Tumor growth \u2192 compressive stress \u2192 displacement",
         "Resection \u2192 stress release \u2192 cavity collapse",
         "Exponential body force models mass-effect release",
         "Max |u|: 0.31 \u2192 2.09 mm with mass effect (6.8\u00d7)",
         "Comparison with ANTs SyN empirical displacement (~4 mm avg)"],
        fig_dir, "fig05_brain_mechanics.png",
        notes="Mass effect modeling is critical for closing the gap "
              "between FEM predictions and empirical observations.")

    _full_figure_slide(prs,
        "FEM vs Empirical Displacement",
        fig_dir, "figC06_01_displacement_comparison.png",
        "Predicted (FEniCSx) vs observed (ANTs SyN) displacement fields",
        notes="Current correlation r\u22480.10 with 2.9\u00d7 magnitude gap. "
              "Mass effect release improves this to r\u22480.31. "
              "Remaining gap: CSF, swelling, tumor mechanics.")


def build_c07(prs, fig_dir):
    """NB_C07: Neural Dynamics \u2014 Hopf Model on the Connectome."""
    _section_divider(
        prs, "Part VII", "Neural Dynamics\nThe Hopf Model on the Connectome",
        "NB_C07 \u2014 Neural mass models, SC modulation, bifurcation shift",
        notes="NB_C07 covers neural dynamics simulation using the "
              "Hopf bifurcation model on the patient connectome.")

    _content_slide(prs,
        "From Neurons to Neural Mass Models",
        ["Population-level dynamics: mean-field approximation",
         "Hopf bifurcation: a\u1d62 controls oscillatory/noisy regime",
         "Each parcel i: d z\u1d62/dt = (a\u1d62 + i\u03c9\u1d62)z\u1d62 "
         "\u2013 |z\u1d62|\u00b2z\u1d62 + \u03b2\u2211 SC\u1d62\u2c7c "
         "(z\u2c7c \u2013 z\u1d62) + \u03c3\u03be\u1d62",
         "Coupling via structural connectivity matrix SC",
         "Simulated FC compared to empirical FC (Pearson)"],
        fig_dir, "figC07_01_hopf_baseline.png",
        notes="The Hopf model produces oscillatory or noise-driven "
              "dynamics depending on the bifurcation parameter a.")

    _content_slide(prs,
        "Mechanics\u2013Dynamics Coupling",
        ["SC modulation: SC\u2032 = SC \u2219 exp(\u2013\u03b1|u|) along tracts",
         "Bifurcation shift: a\u1d62\u2032 = a\u1d62 \u2013 "
         "\u03b3 \u00b7 strain\u1d62",
         "Mechanics \u2192 altered SC \u2192 altered dynamics \u2192 "
         "altered simulated FC",
         "Framework validated: direction correct, magnitude needs "
         "per-region fitting",
         "Coupling parameters: \u03b1 (SC attenuation), \u03b3 "
         "(bifurcation sensitivity)"],
        fig_dir, "figC07_02_coupled_dynamics.png",
        notes="The coupling chain: FEM displacement modulates SC weights "
              "and shifts bifurcation parameters, altering simulated FC.")


def build_c08(prs, fig_dir):
    """NB_C08: Closing the Loop \u2014 Predicted vs Observed."""
    _section_divider(
        prs, "Part VIII", "Closing the Loop\nPredicted vs Observed",
        "NB_C08 \u2014 Cognitive validation: predicted SC disruption vs "
        "\u0394CANTAB",
        notes="NB_C08 closes the loop: do structural changes predicted "
              "by the model correlate with observed cognitive change?")

    _content_slide(prs,
        "The Validation Question",
        ["Can we predict cognitive change from structural disruption?",
         "Map CANTAB domains \u2192 DK68 parcels (literature-based)",
         "Predicted SC disruption \u2192 expected cognitive decrement",
         "Compare with observed \u0394CANTAB z-scores",
         "Quadrant plot: true positives, false positives, "
         "plasticity, confounds"],
        fig_dir, None,
        notes="The quadrant plot reveals where the model succeeds "
              "(correct prediction of decrement) and where it fails "
              "(plasticity, psychosocial confounds).")

    _full_figure_slide(prs,
        "Closing the Loop: Results",
        fig_dir, "figC08_01_closing_the_loop.png",
        "Predicted structural disruption vs observed cognitive change",
        notes="SOC and SSP show correct prediction direction. "
              "RTI shows unexpected improvement (possible practice effect "
              "or compensation).")


def build_c09(prs, fig_dir):
    """NB_C09: Synthesis and Mathematical Perspectives."""
    _section_divider(
        prs, "Part IX", "Synthesis and\nMathematical Perspectives",
        "NB_C09 \u2014 What we demonstrated, what remains, where next",
        notes="NB_C09 synthesizes the proof-of-concept and identifies "
              "the mathematical frontiers.")

    _content_slide(prs,
        "What This Proof-of-Concept Demonstrates",
        ["End-to-end pipeline: raw mpMRI \u2192 digital twin "
         "\u2192 cognitive prediction",
         "Patient-specific geometry, materials, connectivity",
         "Coupled mechanics\u2013dynamics framework: direction correct",
         "Validation against longitudinal observations (unique to glioma)",
         "Open-source, reproducible: NB_C01\u2013C09, GitHub, Colab"],
        fig_dir, None,
        notes="The pipeline is complete end-to-end, though each layer "
              "has quantitative gaps to close.")

    _full_figure_slide(prs,
        "Mathematical Challenges: Roadmap",
        fig_dir, "figC09_01_math_challenges.png",
        "Seven + one challenges and their current status",
        notes="Each challenge maps to specific research directions "
              "and potential collaborations at CAS.")

    _content_slide(prs,
        "From Proof-of-Concept to Translation",
        ["Inverse problem: Bayesian inference on (\u03b8, BC)",
         "Reduced-order models for real-time simulation",
         "Multi-patient validation (25 patients in BTC dataset)",
         "MRE integration for constitutive model validation",
         "Clinical decision support: surgical planning, "
         "outcome prediction"],
        fig_dir, "figC09_02_roadmap.png",
        notes="The path forward requires collaboration between applied "
              "mathematicians, neuroscientists, and clinicians \u2014 "
              "exactly the CAS workshop setting.")


def build_summary(prs, fig_dir):
    """Summary and acknowledgments."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide, CLR_DARK)

    _add_textbox(slide, Inches(0.8), Inches(0.8), Inches(11), Inches(1),
                 "Summary",
                 font_size=44, bold=True, color=CLR_WHITE,
                 font_name="Calibri Light")

    summary_items = [
        ("Glioma data", "Multimodal, longitudinal, single patient "
         "(PAT16) \u2014 pre-op and post-op"),
        ("9-step pipeline", "FreeSurfer \u2192 fMRIPrep \u2192 QSIPrep "
         "\u2192 MRtrix3 \u2192 ANTs \u2192 SC/FC/rsHRF"),
        ("Cavity masking", "Cost-function masking for "
         "topology-changing registration"),
        ("FEM simulation", "FEniCSx, Neo-Hookean + Holzapfel, "
         "2.9 s solve time"),
        ("Mass effect", "Body force model: 0.31 \u2192 2.09 mm "
         "(6.8\u00d7 improvement)"),
        ("Neural dynamics", "Hopf bifurcation on SC, mechanics\u2013"
         "dynamics coupling"),
        ("Cognitive loop", "Predicted SC disruption vs observed "
         "\u0394CANTAB"),
    ]

    paras = []
    for label, desc in summary_items:
        paras.append((f"  {label}:  {desc}", False, 20))

    _add_rich_textbox(slide, Inches(0.8), Inches(2.0), Inches(11.5),
                      Inches(4.5), paras, color=CLR_WHITE,
                      font_name="Calibri Light")

    _set_notes(slide, "Key takeaways. Each point maps to one or more "
               "capstone notebooks.")

    # Acknowledgments
    ack = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(ack, CLR_DARK)

    _add_textbox(ack, Inches(0.8), Inches(1.5), Inches(11), Inches(1),
                 "Acknowledgments",
                 font_size=44, bold=True, color=CLR_WHITE,
                 font_name="Calibri Light")

    _add_rich_textbox(ack, Inches(0.8), Inches(3.0), Inches(11), Inches(3.5),
        [("CAS \u2014 Mathematical Challenges in Brain Mechanics", True, 22),
         ("Henrik Aerts & Daniele Marinazzo (BTC dataset, OpenNeuro)",
          False, 20),
         ("FreeSurfer, fMRIPrep, QSIPrep, MRtrix3, ANTs, FEniCSx teams",
          False, 20),
         ("", False, 12),
         ("github.com/arvidl/glioma-brain-digital-twin-public", False, 18),
         ("github.com/arvidl/BMED365-2026", False, 18)],
        color=CLR_WHITE, font_name="Calibri Light")

    _set_notes(ack, "Thank CAS organizers, data providers, and open-source "
               "tool developers.")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def build_deck(fig_dir: Path, output: Path):
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    build_title_slide(prs)
    build_c01(prs, fig_dir)
    build_c02(prs, fig_dir)
    build_c03(prs, fig_dir)
    build_c04(prs, fig_dir)
    build_c05(prs, fig_dir)
    build_c06(prs, fig_dir)
    build_c07(prs, fig_dir)
    build_c08(prs, fig_dir)
    build_c09(prs, fig_dir)
    build_summary(prs, fig_dir)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    n = len(prs.slides)
    print(f"Saved {n} slides \u2192 {output}")


def main():
    ap = argparse.ArgumentParser(
        description="Build the Glioma Brain Digital Twin PPTX slide deck.")
    ap.add_argument("-o", "--output", type=Path, default=DEFAULT_OUTPUT,
                    help="Output .pptx path (default: %(default)s)")
    ap.add_argument("--fig-dir", type=Path, default=DEFAULT_FIG_DIR,
                    help="Root directory for capstone figures "
                         "(default: %(default)s)")
    args = ap.parse_args()
    build_deck(args.fig_dir, args.output)


if __name__ == "__main__":
    main()
